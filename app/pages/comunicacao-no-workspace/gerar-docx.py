"""Gera o LOGICA-DA-FUNCIONALIDADE.docx a partir do .md da mesma pasta.

O .md é a fonte: edite o .md, rode o lint de escrita e gere o .docx de novo.

Uso (precisa de python-docx):
    python -B gerar-docx.py LOGICA-DA-FUNCIONALIDADE.md LOGICA-DA-FUNCIONALIDADE.docx

Entende título (#, ##, ###), parágrafo em negrito sozinho na linha (vira
subtítulo), lista com "-" e "1.", tabela Markdown, **negrito**, `código` e link.
"""
import re
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ORIGEM, DESTINO = sys.argv[1], sys.argv[2]
linhas = open(ORIGEM, encoding='utf-8').read().split('\n')

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2)

base = doc.styles['Normal']
base.font.name = 'Calibri'
base.font.size = Pt(10.5)
base.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
base.paragraph_format.space_after = Pt(4)
base.paragraph_format.line_spacing = 1.15
for nome, tam in [('Title', 24), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
    st = doc.styles[nome]
    st.font.name = 'Calibri'
    st.font.size = Pt(tam)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
doc.styles['Heading 3'].font.color.rgb = RGBColor(0x40, 0x40, 0x40)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 3'].paragraph_format.space_before = Pt(8)

AZUL = RGBColor(0x1F, 0x5F, 0xBF)


def link(par, url, texto=None):
    parte = par.part
    rid = parte.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    cor = OxmlElement('w:color'); cor.set(qn('w:val'), '1F5FBF'); rpr.append(cor)
    sub = OxmlElement('w:u'); sub.set(qn('w:val'), 'single'); rpr.append(sub)
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = texto or url; t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)
    par._p.append(h)


TOKEN = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|https?://[^\s)|]+)')


def inline(par, texto, negrito=False):
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**'):
            inline(par, parte[2:-2], negrito=True)
        elif parte.startswith('`') and parte.endswith('`'):
            r = par.add_run(parte[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            r.element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            r.bold = negrito
        elif parte.startswith('http'):
            fim = parte.rstrip('.,;')
            link(par, fim)
            if len(fim) < len(parte):
                par.add_run(parte[len(fim):])
        else:
            r = par.add_run(parte)
            r.bold = negrito


def sombra(celula, cor='E8EEF7'):
    tc = celula._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), cor)
    tc.append(s)


def tabela(blocos):
    cab = [c.strip() for c in blocos[0].strip().strip('|').split('|')]
    corpo = [[c.strip() for c in l.strip().strip('|').split('|')] for l in blocos[2:]]
    t = doc.add_table(rows=1 + len(corpo), cols=len(cab))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # o cabeçalho se repete quando a tabela passa de página
    trpr = t.rows[0]._tr.get_or_add_trPr()
    rep = OxmlElement('w:tblHeader'); rep.set(qn('w:val'), 'true'); trpr.append(rep)
    for j, c in enumerate(cab):
        cel = t.rows[0].cells[j]
        cel.text = ''
        inline(cel.paragraphs[0], c, negrito=True)
        sombra(cel)
    for i, linha in enumerate(corpo, start=1):
        for j in range(len(cab)):
            cel = t.rows[i].cells[j]
            cel.text = ''
            inline(cel.paragraphs[0], linha[j] if j < len(linha) else '')
    larguras = None
    if cab[0] == '#':
        larguras = [1.0, 4.6, 3.2, 1.9, 1.9, 4.0]
    elif cab[:3] == ['Dependência', 'Onde se resolve', 'Se faltar']:
        larguras = [6.4, 4.0, 6.2]
    if larguras and len(larguras) == len(cab):
        t.autofit = False
        for row in t.rows:
            for j, cel in enumerate(row.cells):
                cel.width = Cm(larguras[j])
    for row in t.rows:
        for cel in row.cells:
            for p in cel.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    if r.font.name != 'Consolas':
                        r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def item_de_lista(texto, nivel, numero=None):
    p = doc.add_paragraph(style='List Bullet' if numero is None and nivel == 0 else ('List Bullet 2' if numero is None else 'Normal'))
    if numero is not None:
        p.paragraph_format.left_indent = Cm(0.9 + 0.6 * nivel)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.add_run(f'{numero}. ')
    p.paragraph_format.space_after = Pt(2)
    inline(p, texto)


i = 0
paragrafo = []


def fecha_paragrafo():
    global paragrafo
    if paragrafo:
        p = doc.add_paragraph()
        inline(p, ' '.join(paragrafo))
    paragrafo = []


while i < len(linhas):
    l = linhas[i]
    s = l.strip()
    if not s:
        fecha_paragrafo(); i += 1; continue
    if l.startswith('# '):
        fecha_paragrafo(); doc.add_heading(l[2:].strip(), level=0); i += 1; continue
    if l.startswith('## '):
        fecha_paragrafo(); doc.add_heading(l[3:].strip(), level=1); i += 1; continue
    if l.startswith('### '):
        fecha_paragrafo(); doc.add_heading(l[4:].strip(), level=2); i += 1; continue
    if re.fullmatch(r'\*\*[^*]+\*\*', s):
        fecha_paragrafo(); doc.add_heading(s[2:-2], level=3); i += 1; continue
    if s.startswith('|'):
        fecha_paragrafo()
        bloco = []
        while i < len(linhas) and linhas[i].strip().startswith('|'):
            bloco.append(linhas[i]); i += 1
        tabela(bloco); continue
    m = re.match(r'^(\s*)(- |\d+\. )(.*)$', l)
    if m:
        fecha_paragrafo()
        recuo = len(m.group(1)) // 2
        texto = m.group(3)
        i += 1
        # continuação: linhas recuadas que não abrem outro item
        while i < len(linhas) and linhas[i].strip() and re.match(r'^\s{2,}', linhas[i]) and not re.match(r'^\s*(- |\d+\. )', linhas[i]):
            texto += ' ' + linhas[i].strip(); i += 1
        numero = int(m.group(2)[:-2]) if m.group(2)[0].isdigit() else None
        item_de_lista(texto, recuo, numero)
        continue
    paragrafo.append(s)
    i += 1
fecha_paragrafo()

doc.core_properties.title = 'Comunicação no workspace: lógica da funcionalidade'
doc.core_properties.author = 'Produto ENSPACE'
doc.save(DESTINO)
print('ok', DESTINO)

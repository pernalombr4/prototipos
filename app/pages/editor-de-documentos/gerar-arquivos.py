"""
Gera os .docx de verdade que o protótipo abre no Word.

Uso (da raiz do repositório):
    python -B app/pages/editor-de-documentos/gerar-arquivos.py

Saída: app/pages/editor-de-documentos/arquivos/<REFERENCIA>.docx (um por
contrato do mocks.ts) e em-branco.docx. O protótipo importa esses arquivos com
`?url` (Vite), e o botão "Abrir no Word" chama `ms-word:ofe|u|<url>`: o Word do
computador busca o arquivo no endereço do protótipo, como buscaria no ENSPACE.

Os dados abaixo ESPELHAM o mocks.ts (contratante, objeto, valor, vigência).
Mudou o mock, rode de novo. Tudo fictício. Precisa de python-docx.
"""
from pathlib import Path

import docx
from docx.shared import Pt

PASTA = Path(__file__).parent / 'arquivos'

CONTRATOS = [
    ('CTR-0142', 'Aurora Logística Ltda.', 'Consultoria em roteirização de frota', 184500, 12, 4),
    ('CTR-0141', 'Vale Verde Alimentos S.A.', 'Fornecimento de embalagens recicláveis', 62000, 24, 1),
    ('CTR-0139', 'Ponto Norte Engenharia', 'Manutenção predial preventiva', 97300, 12, 2),
    ('CTR-0135', 'Grupo Sereno de Hotelaria', 'Licenciamento de software de reservas', 58900, 12, 2),
    ('CTR-0133', 'Cooperativa Agrícola Serra Alta do Sul de Minas Gerais', 'Assessoria contábil e fiscal', 36000, 12, 7),
    ('CTR-0131', 'Lume Energia Solar', 'Instalação de usinas em telhado', 412000, 18, 1),
    ('CTR-0128', 'Maré Alta Pescados', 'Transporte refrigerado', 128700, 12, 2),
    ('CTR-0126', 'Instituto Raiz de Educação', 'Plataforma de ensino a distância', 75400, 24, 2),
]


def moeda(v: float) -> str:
    inteiro = f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return f'R$ {inteiro}'


def contrato(ref, contratante, objeto, valor, meses, versao):
    d = docx.Document()
    estilo = d.styles['Normal']
    estilo.font.name = 'Cambria'
    estilo.font.size = Pt(11)

    titulo = d.add_paragraph()
    titulo.alignment = 1
    r = titulo.add_run('CONTRATO DE PRESTAÇÃO DE SERVIÇOS')
    r.bold = True
    r.font.size = Pt(14)

    for p in [
        f'CONTRATANTE: {contratante}, pessoa jurídica de direito privado, doravante denominada CONTRATANTE.',
        'CONTRATADA: Jurídico Aurora Serviços Empresariais Ltda., doravante denominada CONTRATADA.',
        f'CLÁUSULA 1. OBJETO. O presente contrato tem por objeto: {objeto.lower()}.',
        f'CLÁUSULA 2. PREÇO. Pelos serviços, a CONTRATANTE pagará o valor total de {moeda(valor)}, em parcelas mensais iguais.',
        f'CLÁUSULA 3. VIGÊNCIA. Este contrato vigora por {meses} meses a partir da assinatura, renovável por igual período mediante termo aditivo.',
        'CLÁUSULA 4. CONFIDENCIALIDADE. As partes manterão sigilo sobre as informações trocadas durante a execução deste contrato.',
        'CLÁUSULA 5. FORO. Fica eleito o foro da comarca de Belo Horizonte para dirimir questões oriundas deste contrato.',
    ]:
        par = d.add_paragraph(p)
        par.alignment = 3

    d.add_paragraph()
    nota = d.add_paragraph(f'Documento fictício do protótipo do ENSPACE. Item {ref}, campo Minuta do contrato, versão {versao}.')
    nota.runs[0].italic = True
    nota.runs[0].font.size = Pt(8)

    # O vínculo com o item, nas propriedades do arquivo (Arquivo › Informações).
    cp = d.core_properties
    cp.title = f'Minuta {ref}'
    cp.subject = f'ENSPACE · {ref} · Minuta do contrato · versão {versao}'
    cp.keywords = f'enspace-item:{ref}; enspace-campo:minuta_do_contrato; enspace-versao:{versao}'
    cp.author = 'ENSPACE (protótipo)'
    return d


def main():
    PASTA.mkdir(exist_ok=True)
    for c in CONTRATOS:
        contrato(*c).save(PASTA / f'{c[0]}.docx')
    branco = docx.Document()
    branco.add_paragraph('')
    branco.core_properties.subject = 'ENSPACE · documento em branco'
    branco.save(PASTA / 'em-branco.docx')
    print(f'{len(CONTRATOS) + 1} arquivos em {PASTA}')


if __name__ == '__main__':
    main()
